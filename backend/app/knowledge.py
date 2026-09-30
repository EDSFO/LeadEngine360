import io
import re
from collections import Counter
from pathlib import Path

from docx import Document as WordDocument
from pypdf import PdfReader


SUPPORTED = {".pdf", ".docx", ".txt", ".md"}
STOP_WORDS = set("a ao aos as o os de da das do dos em no na nos nas por para com sem que e ou um uma seu sua seus suas como mais menos sobre entre pelo pela este esta isso empresa produto serviço software solução".split())


def extract_text(filename: str, raw: bytes) -> str:
    suffix = Path(filename).suffix.lower()
    if suffix == ".pdf":
        reader = PdfReader(io.BytesIO(raw))
        text = "\n".join(page.extract_text() or "" for page in reader.pages)
    elif suffix == ".docx":
        document = WordDocument(io.BytesIO(raw))
        text = "\n".join(paragraph.text for paragraph in document.paragraphs)
        for table in document.tables:
            text += "\n" + "\n".join(" | ".join(cell.text for cell in row.cells) for row in table.rows)
    else:
        text = raw.decode("utf-8-sig", errors="replace")
    text = re.sub(r"\s+", " ", text).strip()
    if len(text) < 40:
        raise ValueError("Não foi possível extrair texto suficiente. PDFs digitalizados precisam de OCR antes do envio.")
    return text


def split_chunks(text: str, size: int = 1100, overlap: int = 160) -> list[str]:
    chunks: list[str] = []
    start = 0
    while start < len(text):
        end = min(len(text), start + size)
        if end < len(text):
            boundary = text.rfind(" ", start + size // 2, end)
            if boundary > start:
                end = boundary
        chunk = text[start:end].strip()
        if chunk:
            chunks.append(chunk)
        if end == len(text):
            break
        start = max(start + 1, end - overlap)
    return chunks


def retrieve(query: str, rows: list[tuple[str, str, str]], limit: int = 8) -> list[dict]:
    query_terms = [word for word in re.findall(r"[\w-]{3,}", query.lower()) if word not in STOP_WORDS]
    query_set = set(query_terms)
    if not query_set:
        query_set = set(re.findall(r"[\w-]{3,}", query.lower()))
    ranked = []
    for chunk_id, filename, content in rows:
        terms = [word for word in re.findall(r"[\w-]{3,}", content.lower()) if word not in STOP_WORDS]
        counts = Counter(terms)
        overlap = sum(min(2, counts[term]) for term in query_set)
        if overlap:
            score = overlap / (len(query_set) ** 0.5 * (len(terms) ** 0.25 or 1))
            ranked.append((score, chunk_id, filename, content))
    ranked.sort(key=lambda row: row[0], reverse=True)
    return [{"chunk_id": chunk_id, "filename": filename, "content": content, "relevance": round(score, 4)} for score, chunk_id, filename, content in ranked[:limit]]
